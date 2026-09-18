import io
import logging
from typing import Dict, List, Optional
from django.core.cache import cache
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import transaction
from django.utils import timezone
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from apps.exchange.base import BaseExchangeDefinition, ErrorCodes, ValidationIssue
from apps.exchange.models import ImportJob
from apps.exchange.registry import exchange_registry
from apps.exchange.services.import_parser import ImportParser, ParsedSheet
from apps.exchange.services.import_validator import ImportValidator, ValidatedRow
from apps.files.models import File
from shared.helpers.cloudinary_service import CloudinaryService

logger = logging.getLogger(__name__)


class ImportService:
    @staticmethod
    def validate_and_preview(job: ImportJob, file_obj, filename: str) -> dict:
        """Parse, validate all layers, generate error report if needed, and prepare preview."""
        definition = exchange_registry.get(job.entity_type)
        if not definition:
            job.status = ImportJob.Status.FAILED
            job.save()
            return {"error": f"Unknown entity type '{job.entity_type}'"}

        job.status = ImportJob.Status.PARSING
        job.file_name = filename
        job.current_step = "Reading spreadsheet..."
        job.progress = 15
        job.save()

        # 1. Parse File
        parsed_sheet = ImportParser.parse_file(file_obj, filename)
        if parsed_sheet.issues:
            job.status = ImportJob.Status.FAILED
            job.error_summary = [i.to_dict() for i in parsed_sheet.issues]
            job.current_step = "File parsing failed"
            job.save()
            return {
                "importId": job.public_id,
                "status": job.status,
                "totalRows": 0,
                "errors": job.error_summary,
            }

        job.current_step = f"Mapping columns for {len(parsed_sheet.rows)} rows..."
        job.progress = 35
        job.save()

        # 2. Structural Layer (Header Mapping)
        header_map, structural_issues = ImportValidator.map_headers(
            parsed_sheet.headers, definition
        )
        has_fatal_structural_error = any(i.severity == "error" for i in structural_issues)
        if has_fatal_structural_error:
            job.status = ImportJob.Status.FAILED
            job.error_summary = [i.to_dict() for i in structural_issues]
            job.total_rows = len(parsed_sheet.rows)
            job.current_step = "Header validation failed"
            job.save()
            return {
                "importId": job.public_id,
                "status": job.status,
                "totalRows": len(parsed_sheet.rows),
                "errors": job.error_summary,
            }

        job.current_step = f"Validating {len(parsed_sheet.rows)} rows..."
        job.progress = 60
        job.save()

        # 3. Field & Business Layers
        validated_rows = ImportValidator.validate_rows(
            raw_rows=parsed_sheet.rows,
            header_map=header_map,
            definition=definition,
            company=job.company,
            mode=job.mode,
            match_by=job.match_by,
        )

        # 4. Aggregate counts
        total_rows = len(validated_rows)
        invalid_rows = sum(1 for r in validated_rows if r.has_errors)
        warning_rows = sum(1 for r in validated_rows if not r.has_errors and r.has_warnings)
        valid_rows = total_rows - invalid_rows

        job.total_rows = total_rows
        job.valid_rows = valid_rows
        job.invalid_rows = invalid_rows
        job.warning_rows = warning_rows

        # 5. Extract error summary
        all_errors = []
        for r in validated_rows:
            for issue in r.issues:
                all_errors.append({
                    "row": r.row_number,
                    "field": issue.field,
                    "code": issue.code,
                    "message": issue.message,
                    "severity": issue.severity,
                    "value": issue.value,
                })
        job.error_summary = all_errors

        # 6. Generate Error Report spreadsheet if errors exist
        if invalid_rows > 0:
            job.current_step = "Generating error report..."
            job.save()
            error_file_bytes = ImportService._generate_error_report(
                parsed_sheet.headers, validated_rows
            )
            error_filename = f"{job.public_id}_error_report.xlsx"
            cache.set(
                f"exchange:import_error_file:{job.public_id}",
                (error_file_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", error_filename),
                timeout=86400,
            )
            file_record = None
            error_report_url = f"/api/v1/exchange/imports/{job.public_id}/error-report/"
            try:
                uploaded_error_file = SimpleUploadedFile(
                    name=error_filename,
                    content=error_file_bytes,
                    content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
                upload_result = CloudinaryService.upload_file(uploaded_error_file, "import_errors")
                if upload_result:
                    file_record = File.update_or_create_file_with_cloudinary(
                        None, upload_result, file_type=File.SYSTEM_TYPE
                    )
                    error_report_url = file_record.get_full_url() or error_report_url
            except Exception as up_exc:
                logger.warning("CloudinaryService error report upload skipped/failed: %s", up_exc)

            job.error_report_file = file_record
            job.error_report_url = error_report_url

        # 7. Cache validated data for commit step
        cache_key = f"exchange:import:{job.public_id}"
        cache_payload = {
            "rows": [
                {
                    "row_number": r.row_number,
                    "data": r.cleaned_data,
                    "action": r.action,
                    "has_errors": r.has_errors,
                    "existing_pk": r.existing_instance.pk if r.existing_instance else None,
                }
                for r in validated_rows
            ]
        }
        cache.set(cache_key, cache_payload, timeout=86400)  # 24h
        job.payload_cache_key = cache_key
        job.status = ImportJob.Status.AWAITING_CONFIRMATION
        job.current_step = "Awaiting user confirmation"
        job.progress = 100
        job.save()

        # 8. Return preview (first 50 rows)
        preview_rows = [r.to_dict() for r in validated_rows[:50]]

        return {
            "importId": job.public_id,
            "status": job.status,
            "totalRows": total_rows,
            "validRows": valid_rows,
            "invalidRows": invalid_rows,
            "warningRows": warning_rows,
            "preview": preview_rows,
            "errorSummary": all_errors[:100],
            "errorReportUrl": job.error_report_url,
        }

    @staticmethod
    def commit_import_job(job: ImportJob) -> ImportJob:
        """Commit validated rows inside an atomic transaction."""
        definition = exchange_registry.get(job.entity_type)
        if not definition:
            job.status = ImportJob.Status.FAILED
            job.save()
            return job

        # Strict atomicity check: if atomic import required and there are invalid rows
        if definition.atomic_import and job.invalid_rows > 0:
            job.status = ImportJob.Status.FAILED
            job.current_step = f"Không thể lưu dữ liệu vì phát hiện {job.invalid_rows} dòng lỗi (Yêu cầu nhập toàn vẹn 100%)."
            job.save()
            return job

        cache_key = job.payload_cache_key or f"exchange:import:{job.public_id}"
        cached_data = cache.get(cache_key)
        if not cached_data or "rows" not in cached_data:
            job.status = ImportJob.Status.FAILED
            job.current_step = "Dữ liệu nhập đã hết hạn trong phiên làm việc. Vui lòng tải lại file."
            job.save()
            return job

        job.status = ImportJob.Status.COMMITTING
        job.started_at = timezone.now()
        job.current_step = "Preparing commit transaction..."
        job.progress = 10
        job.save()

        rows = cached_data["rows"]
        valid_rows = [r for r in rows if not r["has_errors"]]

        created_count = 0
        updated_count = 0
        failed_count = 0

        try:
            with transaction.atomic():
                total_to_process = len(valid_rows)
                for idx, r in enumerate(valid_rows, start=1):
                    row_data = r["data"]
                    action = r["action"]
                    existing_pk = r["existing_pk"]

                    try:
                        if action == "create":
                            definition.create_record(
                                validated_data=row_data,
                                company=job.company,
                                user=job.created_by,
                            )
                            created_count += 1
                        elif action == "update" and existing_pk:
                            existing_instance = definition.model.objects.filter(
                                pk=existing_pk
                            ).first()
                            if existing_instance:
                                definition.update_record(
                                    instance=existing_instance,
                                    validated_data=row_data,
                                    company=job.company,
                                    user=job.created_by,
                                )
                                updated_count += 1
                            else:
                                failed_count += 1
                    except Exception as row_exc:
                        logger.exception("Failed to commit row %d: %s", r["row_number"], row_exc)
                        if definition.atomic_import:
                            raise  # Rollback entire transaction
                        failed_count += 1

                    if idx % 100 == 0 or idx == total_to_process:
                        job.processed_rows = idx
                        job.progress = min(95, 10 + int((idx / max(1, total_to_process)) * 85))
                        job.save()

            job.status = ImportJob.Status.COMPLETED
            job.progress = 100
            job.current_step = "Completed"
            job.completed_at = timezone.now()
            job.created_rows = created_count
            job.updated_rows = updated_count
            job.failed_rows = failed_count
            job.save()

            # Clean up cache
            cache.delete(cache_key)
            logger.info("ImportJob %s committed successfully. Created=%d, Updated=%d", job.public_id, created_count, updated_count)
            return job

        except Exception as exc:
            logger.exception("ImportJob %s transaction failed: %s", job.public_id, exc)
            job.status = ImportJob.Status.FAILED
            job.current_step = f"Lỗi trong quá trình ghi dữ liệu: {str(exc)}"
            job.save()
            return job

    @staticmethod
    def _generate_error_report(headers: List[str], validated_rows: List[ValidatedRow]) -> bytes:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "BaoCaoLoi"

        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        error_row_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
        warning_row_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
        thin_border = Border(
            left=Side(style="thin", color="CBD5E1"),
            right=Side(style="thin", color="CBD5E1"),
            top=Side(style="thin", color="CBD5E1"),
            bottom=Side(style="thin", color="CBD5E1"),
        )

        meta_headers = ["Dòng", "Trạng thái", "Mã lỗi", "Chi tiết lỗi"]
        full_headers = meta_headers + headers
        ws.append(full_headers)

        ws.row_dimensions[1].height = 28
        for col_idx in range(1, len(full_headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border

        for row_idx, r in enumerate(validated_rows, start=2):
            if r.has_errors:
                status_str = "LỖI"
                code_str = "; ".join([i.code for i in r.issues if i.severity == "error"])
                msg_str = "; ".join([f"{i.field}: {i.message}" for i in r.issues if i.severity == "error"])
                row_fill = error_row_fill
            elif r.has_warnings:
                status_str = "CẢNH BÁO"
                code_str = "; ".join([i.code for i in r.issues])
                msg_str = "; ".join([f"{i.field}: {i.message}" for i in r.issues])
                row_fill = warning_row_fill
            else:
                status_str = "HỢP LỆ"
                code_str = ""
                msg_str = ""
                row_fill = None

            row_cells = [
                r.row_number,
                status_str,
                code_str,
                msg_str,
            ] + [r.raw_data.get(h, "") for h in headers]

            ws.append(row_cells)
            ws.row_dimensions[row_idx].height = 20
            for col_idx in range(1, len(row_cells) + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.font = Font(name="Arial", size=10)
                cell.border = thin_border
                cell.alignment = Alignment(vertical="center")
                if row_fill:
                    cell.fill = row_fill

        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(12, min(max_len + 3, 50))

        ws.freeze_panes = "E2"
        out = io.BytesIO()
        wb.save(out)
        out.seek(0)
        return out.getvalue()
