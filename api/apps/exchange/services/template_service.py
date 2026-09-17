import io
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from apps.exchange.base import BaseExchangeDefinition


class TemplateService:
    @staticmethod
    def generate_template(definition: BaseExchangeDefinition) -> bytes:
        """Generate official Excel (.xlsx) template directly from definition contract."""
        wb = openpyxl.Workbook()

        # --- Sheet 1: Template Data Input ---
        ws_data = wb.active
        ws_data.title = "DuLieuNhap"

        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        example_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
        example_font = Font(name="Arial", size=10, italic=True, color="475569")
        thin_border = Border(
            left=Side(style="thin", color="CBD5E1"),
            right=Side(style="thin", color="CBD5E1"),
            top=Side(style="thin", color="CBD5E1"),
            bottom=Side(style="thin", color="CBD5E1"),
        )

        headers = []
        examples = []
        import_field_items = list(definition.import_fields.items())

        for key, field_def in import_field_items:
            req_marker = " *" if field_def.required else ""
            headers.append(f"{field_def.label}{req_marker}")
            examples.append(field_def.example or "")

        ws_data.append(headers)
        ws_data.append(examples)

        # Style header row
        ws_data.row_dimensions[1].height = 30
        for col_idx in range(1, len(headers) + 1):
            cell = ws_data.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = thin_border

        # Style example row
        ws_data.row_dimensions[2].height = 22
        for col_idx in range(1, len(examples) + 1):
            cell = ws_data.cell(row=2, column=col_idx)
            cell.fill = example_fill
            cell.font = example_font
            cell.alignment = Alignment(vertical="center")
            cell.border = thin_border

        # Add Data Validation dropdowns for enum fields
        for col_idx, (key, field_def) in enumerate(import_field_items, start=1):
            if field_def.field_type == "enum" and field_def.choices:
                labels = [str(c[1]).replace('"', '""') for c in field_def.choices]
                formula = f'"{",".join(labels)}"'
                if len(formula) < 255:  # Excel formula length limit
                    dv = DataValidation(
                        type="list",
                        formula1=formula,
                        allow_blank=True,
                        showErrorMessage=True,
                        errorTitle="Giá trị không hợp lệ",
                        error="Vui lòng chọn một giá trị từ danh sách.",
                    )
                    ws_data.add_data_validation(dv)
                    col_letter = get_column_letter(col_idx)
                    dv.add(f"{col_letter}3:{col_letter}1000")

        # Column width calculation
        for col_idx, h in enumerate(headers, start=1):
            col_letter = get_column_letter(col_idx)
            ws_data.column_dimensions[col_letter].width = max(len(h) + 4, 18)

        ws_data.freeze_panes = "A3"

        # --- Sheet 2: Field Instructions & Guide ---
        ws_guide = wb.create_sheet(title="HuongDan")
        ws_guide.views.sheetView[0].showGridLines = True

        guide_headers = ["Tên cột", "Mã trường", "Bắt buộc", "Kiểu dữ liệu", "Giá trị hợp lệ / Ví dụ", "Ghi chú"]
        ws_guide.append(guide_headers)

        guide_header_fill = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
        ws_guide.row_dimensions[1].height = 26
        for col_idx in range(1, len(guide_headers) + 1):
            cell = ws_guide.cell(row=1, column=col_idx)
            cell.fill = guide_header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border

        for row_idx, (key, field_def) in enumerate(import_field_items, start=2):
            choices_str = ", ".join([f"{c[1]} ({c[0]})" for c in field_def.choices]) if field_def.choices else (field_def.example or "")
            ws_guide.append([
                field_def.label,
                key,
                "Bắt buộc" if field_def.required else "Không",
                field_def.field_type,
                choices_str,
                field_def.description or "",
            ])
            for col_idx in range(1, len(guide_headers) + 1):
                cell = ws_guide.cell(row=row_idx, column=col_idx)
                cell.font = Font(name="Arial", size=10)
                cell.border = thin_border
                cell.alignment = Alignment(vertical="center")

        for col in ws_guide.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws_guide.column_dimensions[col_letter].width = max(15, min(max_len + 3, 60))

        out = io.BytesIO()
        wb.save(out)
        out.seek(0)
        return out.getvalue()
